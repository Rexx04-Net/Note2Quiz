import io
import json
import re
import config
from pypdf import PdfReader

try:
    import google.generativeai as genai
    if config.GEMINI_API_KEY:
        genai.configure(api_key=config.GEMINI_API_KEY)
except ImportError:
    genai = None

def extract_text_from_file(file_bytes, filename=""):
    fname = filename.lower() if filename else ""
    
    # 1. PPTX (PowerPoint Presentation XML)
    if fname.endswith(".pptx"):
        try:
            from pptx import Presentation
            prs = Presentation(io.BytesIO(file_bytes))
            text = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        text.append(shape.text)
            extracted = "\n".join(text).strip()
            if extracted:
                return extracted
        except Exception as e:
            print(f"⚠️ [SyllabusParser] PPTX extraction failed: {e}")

    # 2. PDF
    if fname.endswith(".pdf") or not fname or fname.endswith(".pptx"):
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            extracted_text = ""
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    extracted_text += text + "\n"
            if extracted_text.strip():
                return extracted_text.strip()
        except Exception:
            pass

    # 3. Fallback / Plaintext / Legacy PPT text stream scanning
    try:
        utf16_strings = re.findall(rb'(?:[\x20-\x7E\t\r\n]\x00){4,}', file_bytes)
        ascii_strings = re.findall(rb'[\x20-\x7E\t\r\n]{4,}', file_bytes)
        ignore_prefixes = (
            'Current User', 'PowerPoint Document', 'Microsoft', 'Times New Roman',
            'Arial', 'Calibri', 'Courier New', 'Symbol', 'Wingdings', 'Tahoma',
            'Segoe UI', 'SummaryInformation', 'DocumentSummaryInformation'
        )

        extracted_lines = []
        for s in utf16_strings:
            try:
                t = s.decode('utf-16-le', errors='ignore').strip()
                if t and len(t) >= 4 and not t.startswith(ignore_prefixes):
                    t_clean = re.sub(r'[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]', '', t).strip()
                    if len(t_clean) >= 3:
                        extracted_lines.append(t_clean)
            except Exception:
                pass

        for s in ascii_strings:
            try:
                t = s.decode('ascii', errors='ignore').strip()
                if t and len(t) >= 4 and not t.startswith(ignore_prefixes):
                    t_clean = re.sub(r'[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]', '', t).strip()
                    if len(t_clean) >= 3 and not re.match(r'^[0-9a-fA-F]{8,}$', t_clean):
                        extracted_lines.append(t_clean)
            except Exception:
                pass

        seen = set()
        cleaned = []
        for line in extracted_lines:
            if line not in seen:
                seen.add(line)
                cleaned.append(line)

        return "\n".join(cleaned)
    except Exception:
        return ""

def extract_text_from_pdf(pdf_bytes):
    return extract_text_from_file(pdf_bytes, filename="syllabus.pdf")

def generate_fallback_topics(course_name, total_weeks):
    topics = []
    for week in range(1, total_weeks + 1):
        topics.append({
            "week_number": week,
            "topic_title": f"{course_name} - Module {week}",
            "key_concepts": [f"Core concept {week}.1", f"Key principle {week}.2"]
        })
    return topics

def validate_and_normalize_topics(parsed_data, total_weeks, course_name):
    warnings = []
    if not isinstance(parsed_data, dict):
        warnings.append("AI response was not a valid JSON object.")
        return generate_fallback_topics(course_name, total_weeks), warnings, "fallback"

    weekly_topics_raw = parsed_data.get("weekly_topics")
    if not isinstance(weekly_topics_raw, list) or len(weekly_topics_raw) == 0:
        warnings.append("AI output missing 'weekly_topics' list.")
        return generate_fallback_topics(course_name, total_weeks), warnings, "fallback"

    seen_weeks = set()
    cleaned_topics = {}

    for item in weekly_topics_raw:
        if not isinstance(item, dict):
            continue
        
        week_num = item.get("week_number")
        try:
            week_num = int(week_num)
        except (TypeError, ValueError):
            continue

        if week_num < 1 or week_num > total_weeks:
            continue

        if week_num in seen_weeks:
            warnings.append(f"Duplicate week number {week_num} returned by AI; using first occurrence.")
            continue

        seen_weeks.add(week_num)
        
        topic_title = str(item.get("topic_title", "")).strip()
        if not topic_title:
            topic_title = f"{course_name} - Week {week_num} Content"

        key_concepts_raw = item.get("key_concepts", [])
        if isinstance(key_concepts_raw, list):
            key_concepts = [str(kc).strip() for kc in key_concepts_raw if str(kc).strip()]
        else:
            key_concepts = [str(key_concepts_raw).strip()]

        if not key_concepts:
            key_concepts = [topic_title]

        cleaned_topics[week_num] = {
            "week_number": week_num,
            "topic_title": topic_title,
            "key_concepts": key_concepts
        }

    # Normalize missing weeks
    final_topics = []
    for w in range(1, total_weeks + 1):
        if w in cleaned_topics:
            final_topics.append(cleaned_topics[w])
        else:
            warnings.append(f"Week {w} missing from AI output; padded with default topic.")
            final_topics.append({
                "week_number": w,
                "topic_title": f"{course_name} - Week {w} Self-Study & Revision",
                "key_concepts": ["Review week material", "Practice problem set"]
            })

    source = "gemini" if len(cleaned_topics) > 0 else "fallback"
    return final_topics, warnings, source

def parse_syllabus_pdf(pdf_bytes, total_weeks, course_name, filename=""):
    warnings = []
    pdf_text = extract_text_from_file(pdf_bytes, filename=filename)

    if not config.GEMINI_API_KEY or genai is None:
        warnings.append("Gemini API key is not configured. Generated deterministic fallback schedule.")
        return generate_fallback_topics(course_name, total_weeks), warnings, "fallback"

    prompt = f"""
You are an expert syllabus parsing assistant for an educational application.
Analyze the following course syllabus and teaching plan text for the course "{course_name}".
Extract exactly {total_weeks} weekly lecture topics for a {total_weeks}-week semester.

RULES:
1. Return ONLY valid, raw JSON with NO markdown formatting, NO code blocks (do NOT wrap in ```json).
2. Follow this exact JSON schema:
{{
    "course_name": "{course_name}",
    "weekly_topics": [
        {{
            "week_number": 1,
            "topic_title": "Topic_0_Introduction-The_XR_Mission",
            "key_concepts": ["Concept 1", "Concept 2"]
        }}
    ]
}}
3. Include week_number from 1 up to {total_weeks}.
4. TOPIC & FILE NAME NOMENCLATURE (CRITICAL):
   - You MUST extract and preserve the exact topic codes, file names, or chapter numbers from the syllabus (e.g., "Topic_0_Introduction-The_XR_Mission", "Topic_1_Introduction_to_Metaverse", "Topic_2_Fundamentals_of_XR_and_Unity_Introduction").
   - If a week covers multiple topics (e.g. Topic 0 & Topic 1 in Week 1), include them with their topic numbers/names.
   - Do NOT omit the Topic numbers (Topic 0, Topic 1, Topic 2, etc.) or lecture file titles.
5. If the syllabus lacks specific details for a certain week, generate a reasonable topic title following this naming format.

Syllabus Text:
{pdf_text[:20000]}
"""

    models_to_try = [
        'gemini-2.5-flash',
        'models/gemini-2.5-flash',
        'gemini-1.5-flash',
        'gemini-2.0-flash',
    ]

    raw_response = None
    for model_name in models_to_try:
        try:
            model = genai.GenerativeModel(model_name)
            resp = model.generate_content(prompt)
            if resp and resp.text:
                raw_response = resp.text
                break
        except Exception as e:
            print(f"⚠️ [SyllabusParser] Model {model_name} failed: {e}")
            continue

    if not raw_response:
        warnings.append("Gemini API request failed for all model targets. Used fallback syllabus breakdown.")
        return generate_fallback_topics(course_name, total_weeks), warnings, "fallback"

    # Clean response code fences if present
    cleaned_json = raw_response.strip()
    cleaned_json = re.sub(r"^```(?:json)?\s*", "", cleaned_json)
    cleaned_json = re.sub(r"\s*```$", "", cleaned_json)

    try:
        parsed = json.loads(cleaned_json)
        topics, norm_warnings, source = validate_and_normalize_topics(parsed, total_weeks, course_name)
        warnings.extend(norm_warnings)
        return topics, warnings, source
    except Exception as parse_err:
        warnings.append(f"Failed to parse JSON output from Gemini ({parse_err}). Used fallback schedule.")
        return generate_fallback_topics(course_name, total_weeks), warnings, "fallback"
